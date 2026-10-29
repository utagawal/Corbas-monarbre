"""Tests de non-régression : python -m unittest discover -s tests -v   (depuis le dossier parrainage)."""
import atexit
import html
import io
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import time
import unittest
import zipfile
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

TMP = tempfile.mkdtemp(prefix="parrainage-test-")
atexit.register(shutil.rmtree, TMP, True)
os.environ.update(DATA_DIR=TMP, ADMIN_EMAIL="admin@test.local", ADMIN_PASSWORD="motdepasse-test-1", DEMO="0", HTTPS="0")
for k in ("SMTP_HOST", "SECRET_KEY", "TRUST_PROXY"):
    os.environ.pop(k, None)
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app  # noqa: E402
import certificat  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from werkzeug.security import check_password_hash, generate_password_hash  # noqa: E402


def _hors_ligne(*a, **k):
    raise OSError("hors ligne (test)")


certificat.requests.get = _hors_ligne  # pas d'accès réseau pendant les tests : la carte utilise son fond de secours
PDF = (b"%PDF-1.4\n%test\n", "livret.pdf")
MDP = "motdepasse-test-1"


class Base(unittest.TestCase):
    def setUp(self):
        app._essais.clear()
        app._cache_api.clear()
        os.environ.pop("NOTIFICATION_EMAILS", None)
        os.environ.pop("APP_URL", None)
        app.LIMITE_ENVOIS = 100
        app.LIMITE_API = 90
        app.AUTOCOMPLETION = True
        con = sqlite3.connect(app.DB_PATH)
        for t in ("parrainages", "arbres", "essences", "campagnes", "mails", "users", "parametres"):
            con.execute(f"DELETE FROM {t}")
        con.execute("INSERT INTO users(email,hash,role) VALUES('admin@test.local',?,'admin')", (generate_password_hash(MDP),))
        con.execute("INSERT INTO users(email,hash,role) VALUES('agent@test.local',?,'agent')", (generate_password_hash(MDP),))
        con.execute("INSERT INTO campagnes(id,nom,annee,ouverture,fermeture,ouverte) VALUES(1,'Test 2026',2026,'2000-01-01','2100-01-01',1)")
        for i, (nom, actif) in enumerate((("Chêne", 1), ("Tilleul", 1), ("Orme", 0)), 1):
            con.execute("INSERT INTO essences(id,nom,latin,resume,symbole,caracteristiques,role,couleur,actif) VALUES(?,?,?,?,?,?,?,?,?)",
                        (i, nom, "Latin " + nom, "Résumé " + nom, "Symbole " + nom, "Caractéristiques " + nom, "Rôle " + nom, "#2e6b34", actif))
        for n, ess in enumerate((1, 1, 1, 2, 2, 3), 1):
            con.execute("INSERT INTO arbres(code,campagne_id,essence_id,lieu,adresse,lat,lon,annee_plantation) VALUES(?,?,?,?,?,?,?,?)",
                        (f"COR-2026-{n:03d}", 1, ess, "Square test", "Rue test", 45.6706 + n * 0.0001, 4.8960, 2026))
        con.commit()
        con.close()
        for f in os.listdir(app.LIVRETS):
            os.remove(os.path.join(app.LIVRETS, f))
        for f in os.listdir(app.PHOTOS_ENVOYEES):
            os.remove(os.path.join(app.PHOTOS_ENVOYEES, f))
        self.pub = app.app.test_client()

    # ---- outils
    def sql(self, req, *a):
        con = sqlite3.connect(app.DB_PATH)
        con.row_factory = sqlite3.Row
        rows = con.execute(req, a).fetchall()
        con.commit()
        con.close()
        return rows

    def un(self, req, *a):
        r = self.sql(req, *a)
        return r[0] if r else None

    def donnees(self, **kw):
        n = app.aujourdhui()
        d = dict(enfant_prenom="Léa", enfant_nom="Martin", naissance=(n - timedelta(days=60)).isoformat(),
                 lien="Domicilié(e) à Corbas", parent_qualite="Mère", parent_nom="Julie Martin", email="julie@example.org",
                 email2="julie@example.org", tel="06 12 34 56 78", adresse="3 rue des Lilas", cp="69960", ville="Corbas",
                 essence1="1", essence2="2", attest="1", engage="1", rgpd="1")
        d.update(kw)
        return {k: v for k, v in d.items() if v is not None}

    def envoyer(self, fichier=PDF, client=None, **kw):
        d = self.donnees(**kw)
        if fichier is not None:
            d["livret"] = (io.BytesIO(fichier[0]), fichier[1])
        return (client or self.pub).post("/demande", data=d, content_type="multipart/form-data")

    def cree(self, **kw):
        r = self.envoyer(**kw)
        self.assertEqual(r.status_code, 302, r.get_data(as_text=True)[:3000])
        return self.un("SELECT id FROM parrainages ORDER BY id DESC")["id"]

    def connecte(self, email="admin@test.local", mdp=MDP):
        c = app.app.test_client()
        r = c.post("/admin/login", data=dict(email=email, password=mdp))
        self.assertEqual(r.status_code, 302)
        return c

    def jeton(self, c):
        return re.search(r'name="csrf" value="([0-9a-f]+)"', c.get("/admin/mot-de-passe").get_data(as_text=True)).group(1)

    def agir(self, c, pid, action, **kw):
        r = c.post(f"/admin/demande/{pid}/action", data=dict(csrf=self.jeton(c), action=action, **kw), follow_redirects=True)
        return html.unescape(r.get_data(as_text=True))

    def poster(self, c, url, **kw):
        return html.unescape(c.post(url, data=dict(csrf=self.jeton(c), **kw), follow_redirects=True).get_data(as_text=True))


class TestFormulairePublic(Base):
    def test_affichage(self):
        r = self.pub.get("/")
        self.assertEqual(r.status_code, 200)
        html = r.get_data(as_text=True)
        self.assertIn("Chêne", html)
        self.assertIn("Tilleul", html)
        self.assertNotIn("Orme", html)  # essence désactivée : jamais proposée
        self.assertIn('maxlength="60"', html)

    def test_demande_valide(self):
        r = self.envoyer()
        self.assertEqual(r.status_code, 302)
        self.assertIn("/merci/", r.headers["Location"])
        p = self.un("SELECT * FROM parrainages")
        self.assertEqual((p["ref"], p["statut"], p["enfant_prenom"], p["essence2_id"]), (f"PAR-{app.aujourdhui().year}-0001", "Reçue", "Léa", 2))
        self.assertTrue(p["livret"].endswith(".pdf"))
        self.assertTrue(os.path.exists(os.path.join(app.LIVRETS, p["livret"])))
        sujets = [m["sujet"] for m in self.sql("SELECT sujet FROM mails")]
        self.assertTrue(any("Bienvenue à Léa" in s for s in sujets))
        self.assertEqual(len([s for s in sujets if s.startswith("Nouvelle demande")]), 2)  # un courrier par compte
        page = self.pub.get(r.headers["Location"]).get_data(as_text=True)
        self.assertIn("Merci <strong>Julie</strong>", page)
        self.assertIn("Cache-Control", "".join(self.pub.get(r.headers["Location"]).headers.keys()))

    def test_aides_a_la_saisie(self):
        h = self.pub.get("/").get_data(as_text=True)
        self.assertIn("Avant de commencer", h)
        self.assertIn("champs marqués d'un *", h)
        self.assertIn(f'min="{app.mois_avant(app.aujourdhui(), 12).isoformat()}"', h)  # le sélecteur de date grise les dates trop anciennes
        self.assertIn('placeholder="prenom.nom@exemple.fr"', h)
        self.assertNotIn("Des arbres sont encore disponibles", h)  # bandeau inutile quand le stock est confortable
        self.sql("UPDATE arbres SET statut='parraine' WHERE id>2")
        self.assertIn("Il reste 2 arbres disponibles", self.pub.get("/").get_data(as_text=True))

    def test_page_confirmation(self):
        self.cree()
        ref = self.un("SELECT ref FROM parrainages")["ref"]
        h = self.pub.get("/merci/" + app._jeton.dumps(ref)).get_data(as_text=True)
        self.assertIn("<strong>maintenant</strong>", h)
        self.assertIn(ref, h)
        self.assertIn("original du livret de famille", h)
        self.assertNotIn("liste d'attente", h)
        self.sql("UPDATE parrainages SET statut=\"Liste d'attente\"")
        self.assertIn("Vous êtes sur liste d'attente", self.pub.get("/merci/" + app._jeton.dumps(ref)).get_data(as_text=True))

    def test_page_merci_non_devinable(self):
        self.cree()
        ref = self.un("SELECT ref FROM parrainages")["ref"]
        self.assertEqual(self.pub.get(f"/merci/{ref}").status_code, 404)
        jeton = app._jeton.dumps(ref)
        self.assertEqual(self.pub.get(f"/merci/{jeton}").status_code, 200)
        self.assertEqual(self.pub.get(f"/merci/{jeton[:-3]}abc").status_code, 404)

    def test_age_limite(self):
        n = app.aujourdhui()
        for delta, ok in ((app.mois_avant(n, 12), True), (app.mois_avant(n, 12) - timedelta(days=1), False),
                          (n, True), (n + timedelta(days=1), False)):
            app._essais.clear()
            nom = f"Test{delta.isoformat()}"
            r = self.envoyer(naissance=delta.isoformat(), enfant_prenom=nom)
            self.assertEqual(r.status_code == 302, ok, delta)
        r = self.envoyer(naissance="n'importe quoi")
        self.assertEqual(r.status_code, 400)
        self.assertIn("date de naissance", r.get_data(as_text=True))

    def test_mois_avant(self):
        from datetime import date
        self.assertEqual(app.mois_avant(date(2024, 2, 29), 12), date(2023, 2, 28))
        self.assertEqual(app.mois_avant(date(2026, 3, 31), 1), date(2026, 2, 28))
        self.assertEqual(app.mois_avant(date(2026, 1, 15), 12), date(2025, 1, 15))
        self.assertEqual(app.mois_avant(date(2026, 10, 29), 12), date(2025, 10, 29))

    def test_domicile(self):
        for i, (ville, cp, lien, ok) in enumerate((("Lyon", "69003", "Domicilié(e) à Corbas", False), ("CORBAS", "69960", "Domicilié(e) à Corbas", True),
                                    (" corbas ", "69960", "Domicilié(e) à Corbas", True), ("Corbas", "69000", "Domicilié(e) à Corbas", False),
                                    ("Lyon", "69003", "Né(e) à Corbas", True), ("Lyon", "69003", "Né et domicilié à Corbas", False),
                                    ("Corbas", "6996", "Né(e) à Corbas", False))):
            app._essais.clear()
            r = self.envoyer(ville=ville, cp=cp, lien=lien, enfant_prenom=f"Cas{i}")
            self.assertEqual(r.status_code == 302, ok, (ville, cp, lien, r.get_data(as_text=True)[:200]))

    def test_courriel_et_telephone(self):
        for kw in (dict(email="pas-un-mail", email2="pas-un-mail"), dict(email2="autre@example.org"), dict(email="a b@example.org", email2="a b@example.org"),
                   dict(tel="appelez-moi"), dict(parent_qualite="Oncle"), dict(lien="Autre")):
            r = self.envoyer(**kw)
            self.assertEqual(r.status_code, 400, kw)
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 0)

    def test_telephone_mobile_obligatoire(self):
        r = self.envoyer(tel="")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Indiquez votre numéro de mobile", r.get_data(as_text=True))
        for faux in ("01 23 45 67 89", "+33 1 23 45 67 89", "0612", "appelez-moi", "06 12 34 56 789", "+123"):
            self.assertEqual(self.envoyer(tel=faux).status_code, 400, faux)
        self.assertIn("(06 ou 07", self.envoyer(tel="01 23 45 67 89").get_data(as_text=True))
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 0)
        # formes acceptées : toutes sont enregistrées sous la même forme lisible
        for i, (saisi, attendu) in enumerate((("0612345678", "06 12 34 56 78"), ("06.12.34.56.78", "06 12 34 56 78"),
                                              ("+33 6 12 34 56 78", "06 12 34 56 78"), ("+33 (0)7 12 34 56 78", "07 12 34 56 78"),
                                              ("0033612345678", "06 12 34 56 78"), ("+32 470 12 34 56", "+32470123456"))):
            app._essais.clear()
            self.cree(tel=saisi, enfant_prenom=f"Tel{i}")
            self.assertEqual(self.un("SELECT tel FROM parrainages ORDER BY id DESC")["tel"], attendu, saisi)
            self.sql("UPDATE parrainages SET statut='Annulée'")  # libère la place réservée (stock limité à 3 chênes)
        h = self.pub.get("/").get_data(as_text=True)
        self.assertIn("Téléphone mobile *", h)
        self.assertRegex(h, r'<input type="tel"[^>]*required')

    def test_champs_obligatoires_et_longueurs(self):
        for k in ("enfant_prenom", "enfant_nom", "parent_nom", "adresse", "ville", "tel", "attest", "engage", "rgpd", "essence1"):
            r = self.envoyer(**{k: ""})
            self.assertEqual(r.status_code, 400, k)
        self.assertEqual(self.envoyer(enfant_prenom="x" * 61).status_code, 400)
        self.assertEqual(self.envoyer(adresse="x" * 151).status_code, 400)
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 0)

    def test_normalisation(self):
        self.cree(enfant_prenom="  léa\n marie ", enfant_nom="De   la\tTour")
        p = self.un("SELECT * FROM parrainages")
        self.assertEqual((p["enfant_prenom"], p["enfant_nom"]), ("Léa Marie", "De la Tour"))
        self.assertEqual(app.soigner_prenom("JEAN-PIERRE"), "Jean-Pierre")
        self.assertEqual(app.soigner_prenom("Jean-Pierre"), "Jean-Pierre")
        self.assertEqual(app.soigner_prenom("McKenzie"), "McKenzie")  # saisie mixte : respectée

    def test_livret(self):
        self.assertEqual(self.envoyer(fichier=None).status_code, 400)
        self.assertEqual(self.envoyer(fichier=(b"juste du texte", "livret.pdf")).status_code, 400)  # contenu non reconnu malgré l'extension
        self.assertEqual(self.envoyer(fichier=(b"", "vide.pdf")).status_code, 400)
        self.assertEqual(self.envoyer(fichier=(b"%PDF" + b"0" * (5 * 1024 * 1024), "gros.pdf")).status_code, 400)
        self.assertEqual(self.envoyer(fichier=(b"%PDF" + b"0" * (6 * 1024 * 1024), "enorme.pdf")).status_code, 413)
        self.assertEqual(os.listdir(app.LIVRETS), [])  # aucun fichier orphelin après les refus
        for contenu, ext in ((bytes([0xFF, 0xD8, 0xFF, 0xE0]) + b"0" * 20, ".jpg"), (bytes([0x89, 0x50, 0x4E, 0x47]) + b"0" * 20, ".png")):
            app._essais.clear()
            self.cree(fichier=(contenu, "photo.bin"), enfant_prenom="Img" + ext.strip("."))
            self.assertTrue(self.un("SELECT livret FROM parrainages ORDER BY id DESC")["livret"].endswith(ext))

    def test_aucun_fichier_orphelin_si_refus(self):
        self.assertEqual(self.envoyer(attest=None).status_code, 400)
        self.assertEqual(os.listdir(app.LIVRETS), [])

    def test_doublon_et_jumeaux(self):
        self.cree()
        r = self.envoyer(enfant_prenom="LÉA", enfant_nom="martin")  # accents et casse ignorés
        self.assertEqual(r.status_code, 400)
        self.assertIn("existe déjà", r.get_data(as_text=True))
        self.cree(enfant_prenom="Zoé")  # jumeau : prénom différent, même date de naissance
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 2)
        # une demande refusée ne bloque pas un nouveau dépôt
        c = self.connecte()
        pid = self.un("SELECT id FROM parrainages WHERE enfant_prenom='Léa'")["id"]
        self.agir(c, pid, "refuser", motif="Pièce illisible")
        self.cree()
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 3)

    def test_doublon_prenom_accentue(self):
        self.cree(enfant_prenom="Éloïse")
        self.assertEqual(self.envoyer(enfant_prenom="ELOISE").status_code, 400)

    def test_pot_de_miel(self):
        r = self.envoyer(site_web="http://spam.example")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 0)

    def test_limitation_des_envois(self):
        app.LIMITE_ENVOIS = 2
        for i in range(2):
            self.cree(enfant_prenom=f"Enfant{i}")
        r = self.envoyer(enfant_prenom="Trop")
        self.assertEqual(r.status_code, 429)
        self.assertIn("Trop de demandes", r.get_data(as_text=True))

    def test_les_erreurs_ne_consomment_pas_le_quota(self):
        app.LIMITE_ENVOIS = 2
        for _ in range(5):
            self.assertEqual(self.envoyer(attest=None).status_code, 400)
        self.cree(enfant_prenom="A")
        self.cree(enfant_prenom="B")

    def test_essence_non_proposee(self):
        self.assertEqual(self.envoyer(essence1="3").status_code, 400)  # essence désactivée
        self.assertEqual(self.envoyer(essence1="99").status_code, 400)
        self.cree(essence2="3", enfant_prenom="Sans2")
        self.assertIsNone(self.un("SELECT essence2_id FROM parrainages")["essence2_id"])  # seconde essence invalide : ignorée

    def test_stock_net_liste_attente_automatique(self):
        # 3 chênes + 2 tilleuls proposables (l'orme est désactivé) : 5 places
        for i in range(5):
            self.cree(enfant_prenom=f"E{i}", essence1="1" if i < 3 else "2", essence2=None)
        self.assertIn("attribuées ou en cours d'attribution", self.pub.get("/").get_data(as_text=True))
        self.cree(enfant_prenom="Attente", essence1="1", essence2=None)
        self.assertEqual(self.un("SELECT statut FROM parrainages ORDER BY id DESC")["statut"], "Liste d'attente")
        mails = " ".join(m["corps"] for m in self.sql("SELECT corps FROM mails WHERE destinataire='julie@example.org' ORDER BY id DESC LIMIT 1"))
        self.assertIn("liste d'attente", mails)
        # une demande refusée libère une place
        c = self.connecte()
        self.agir(c, 1, "refuser", motif="Hors critères")
        self.assertNotIn("attribuées ou en cours d'attribution", self.pub.get("/").get_data(as_text=True))
        self.cree(enfant_prenom="Nouvelle", essence1="1", essence2=None)
        self.assertEqual(self.un("SELECT statut FROM parrainages ORDER BY id DESC")["statut"], "Reçue")

    def test_stock_net_par_essence(self):
        for i in range(3):
            self.cree(enfant_prenom=f"C{i}", essence1="1", essence2=None)
        page = self.pub.get("/").get_data(as_text=True)
        self.assertNotIn('name="essence1" value="1"', page)  # plus de chêne proposé
        self.assertIn('name="essence1" value="2"', page)
        self.assertEqual(self.envoyer(enfant_prenom="Tard", essence1="1").status_code, 400)
        self.cree(enfant_prenom="Tilleul", essence1="2", essence2=None)

    def test_essence_desactivee_ne_compte_pas(self):
        self.sql("UPDATE arbres SET statut='parraine' WHERE essence_id IN (1,2)")  # seul l'orme (inactif) reste libre
        self.cree(enfant_prenom="Attente", essence1="1", essence2=None)
        self.assertEqual(self.un("SELECT statut FROM parrainages")["statut"], "Liste d'attente")

    def test_campagne_fermee(self):
        for sql in ("UPDATE campagnes SET ouverte=0", "UPDATE campagnes SET ouverte=1, ouverture='2999-01-01'",
                    "UPDATE campagnes SET ouverte=1, ouverture='2000-01-01', fermeture='2001-01-01'"):
            self.sql(sql)
            self.assertIn("fermées", self.pub.get("/").get_data(as_text=True))
            self.envoyer()
            self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 0)

    def test_en_tetes(self):
        r = self.pub.get("/")
        csp = r.headers["Content-Security-Policy"]
        self.assertIn("frame-ancestors 'self' https://corbas.fr https://www.corbas.fr https://ville-corbas.fr", csp)
        self.assertIn("form-action 'self'", csp)
        self.assertEqual(r.headers["X-Content-Type-Options"], "nosniff")
        self.assertNotIn("X-Frame-Options", r.headers)
        a = self.connecte().get("/admin/")
        self.assertEqual(a.headers["X-Frame-Options"], "DENY")
        self.assertIn("frame-ancestors 'none'", a.headers["Content-Security-Policy"])
        self.assertEqual(a.headers["Cache-Control"], "no-store")

    def test_ancetres_sans_apostrophes(self):
        self.assertEqual(app.ancetres("self https://a.fr"), "'self' https://a.fr")
        self.assertEqual(app.ancetres("'self' https://a.fr, https://b.fr"), "'self' https://a.fr https://b.fr")
        self.assertEqual(app.ancetres(""), "'self'")

    def test_sante_et_erreurs(self):
        self.assertEqual(self.pub.get("/sante").get_data(as_text=True), "ok")
        r = self.pub.get("/merci/n-importe-quoi")
        self.assertEqual(r.status_code, 404)
        self.assertIn("Page introuvable", r.get_data(as_text=True))

    def test_reference_en_cas_de_collision(self):
        self.cree()
        vraie = app.ref_suivante
        appels = []

        def capricieuse():
            appels.append(1)
            return self.un("SELECT ref FROM parrainages")["ref"] if len(appels) == 1 else vraie()
        app.ref_suivante = capricieuse
        try:
            self.cree(enfant_prenom="Autre")
        finally:
            app.ref_suivante = vraie
        self.assertEqual(len(appels), 2)
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages")["n"], 2)

    def test_heure_de_paris(self):
        self.assertEqual(app.aujourdhui(), datetime.now(ZoneInfo("Europe/Paris")).date())


class TestSuggestionsAdresses(Base):
    """Les suggestions passent par notre serveur (relais) : on remplace l'appel au service extérieur par un faux."""

    def setUp(self):
        super().setUp()
        self.appels = []
        self.vrai = app.appel_json
        app.appel_json = self.faux

    def tearDown(self):
        app.appel_json = self.vrai

    def faux(self, url, params):
        self.appels.append((url, dict(params)))
        if "panne" in params["q"]:
            raise ConnectionError("service injoignable")
        if "photon" in url:
            return {"features": [
                {"geometry": {"coordinates": [4.902782, 45.665988]}, "properties": {"name": "École primaire Jean Jaurès", "street": "Chemin de Grange Blanche", "postcode": "69960", "city": "Corbas"}},
                {"geometry": {"coordinates": [4.9, 45.67]}, "properties": {"housenumber": "3", "street": "Rue des Lilas", "postcode": "69960", "city": "Corbas"}},
                {"geometry": {"coordinates": [4.9, 45.67]}, "properties": {"housenumber": "3", "street": "Rue des Lilas", "postcode": "69960", "city": "Corbas"}},
                {"geometry": {"coordinates": [4.9, 45.67]}, "properties": {"postcode": "69960", "city": "Corbas"}}]}
        return {"features": [
            {"properties": {"type": "housenumber", "name": "3 Rue des Lilas", "label": "3 Rue des Lilas 69960 Corbas", "postcode": "69960", "city": "Corbas"}},
            {"properties": {"type": "street", "name": "Rue des Lilas", "label": "Rue des Lilas 69960 Corbas", "postcode": "69960", "city": "Corbas"}},
            {"properties": {"type": "municipality", "name": "Corbas", "label": "Corbas", "postcode": "69960", "city": "Corbas"}}]}

    def test_adresses_publiques(self):
        r = self.pub.get("/api/adresse?q=3 rue des lilas")
        self.assertEqual(r.status_code, 200)
        res = r.get_json()["resultats"]
        self.assertEqual([x["adresse"] for x in res], ["3 Rue des Lilas", "Rue des Lilas"])  # la commune seule n'est pas une adresse
        self.assertEqual((res[0]["cp"], res[0]["ville"], res[0]["numero"], res[1]["numero"]), ("69960", "Corbas", True, False))
        url, params = self.appels[0]
        self.assertIn("geopf", url)
        self.assertNotIn("citycode", params)  # enfant né à Corbas : adresses de toute la France
        self.pub.get("/api/adresse?q=rue des&corbas=1")
        self.assertEqual(self.appels[1][1]["citycode"], "69273")  # enfant domicilié à Corbas : adresses de Corbas seulement

    def test_requete_trop_courte_et_cache(self):
        self.assertEqual(self.pub.get("/api/adresse?q=ab").get_json(), {"resultats": []})
        self.assertEqual(self.pub.get("/api/adresse").get_json(), {"resultats": []})
        self.assertEqual(self.appels, [])  # rien n'est envoyé au service extérieur
        for _ in range(3):
            self.pub.get("/api/adresse?q=3 Rue des Lilas")
        self.pub.get("/api/adresse?q=3 rue DES lilas")  # casse ignorée
        self.assertEqual(len(self.appels), 1)  # réponses mises en cache

    def test_service_en_panne(self):
        r = self.pub.get("/api/adresse?q=panne de service")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json(), {"resultats": [], "indisponible": True})  # le formulaire continue de fonctionner
        self.pub.get("/api/adresse?q=panne de service")
        self.assertEqual(len(self.appels), 2)  # une panne n'est jamais mise en cache

    def test_limitation_et_desactivation(self):
        app.LIMITE_API = 3
        for i in range(3):
            self.assertEqual(self.pub.get(f"/api/adresse?q=rue numero {i}").status_code, 200)
        r = self.pub.get("/api/adresse?q=rue numero 9")
        self.assertEqual((r.status_code, r.get_json()["indisponible"]), (429, True))
        app._essais.clear()
        app.AUTOCOMPLETION = False
        self.assertEqual(self.pub.get("/api/adresse?q=3 rue des lilas").get_json(), {"resultats": [], "indisponible": True})

    def test_lieux_pour_les_arbres(self):
        self.assertEqual(self.pub.get("/admin/api/lieu?q=ecole").status_code, 302)  # réservé aux administrateurs connectés
        self.assertEqual(self.connecte("agent@test.local").get("/admin/api/lieu?q=ecole jean").status_code, 403)
        r = self.connecte().get("/admin/api/lieu?q=ecole jean jaures")
        res = r.get_json()["resultats"]
        self.assertEqual(len(res), 2)  # doublon écarté, entrée sans nom ni rue écartée
        self.assertEqual((res[0]["lieu"], res[0]["adresse"], res[0]["lat"], res[0]["lon"]),
                         ("École primaire Jean Jaurès", "Chemin de Grange Blanche, 69960 Corbas", 45.665988, 4.902782))
        self.assertEqual(res[1]["lieu"], "3 Rue des Lilas")
        url, params = self.appels[0]
        self.assertIn("photon", url)
        self.assertEqual(params["bbox"], app.BBOX_CORBAS)

    def test_pages_chargent_la_saisie_assistee(self):
        h = self.pub.get("/").get_data(as_text=True)
        self.assertIn("adresse.js", h)
        self.assertIn("/api/adresse", h)
        self.assertIn("Base Adresse Nationale", h)  # mention d'information dans le texte RGPD
        a = self.connecte().get("/admin/arbres").get_data(as_text=True)
        self.assertIn("/admin/api/lieu", a)
        self.assertIn("Rechercher un lieu ou une adresse", a)
        js = self.pub.get("/static/adresse.js")
        self.assertEqual(js.status_code, 200)
        self.assertIn("window.Autocompletion", js.get_data(as_text=True))


class TestNotificationsAgents(Base):
    """Un e-mail prévient les agents à chaque nouvelle demande : adresses réglables (plusieurs), sans donnée personnelle."""

    def notifs(self):
        return self.sql("SELECT destinataire, sujet, corps FROM mails WHERE sujet LIKE 'Nouvelle demande%' ORDER BY id")

    def test_par_defaut_tous_les_comptes(self):
        self.cree()
        self.assertEqual(sorted(m["destinataire"] for m in self.notifs()), ["admin@test.local", "agent@test.local"])

    def test_plusieurs_adresses_du_serveur(self):
        os.environ["NOTIFICATION_EMAILS"] = "Accueil@Mairie.fr, urbanisme@mairie.fr;  accueil@mairie.fr\nenfance@mairie.fr  adresse-invalide"
        try:
            self.cree()
        finally:
            os.environ.pop("NOTIFICATION_EMAILS")
        # minuscules, sans doublon, adresse incorrecte ignorée, et les comptes ne sont plus prévenus
        self.assertEqual([m["destinataire"] for m in self.notifs()], ["accueil@mairie.fr", "urbanisme@mairie.fr", "enfance@mairie.fr"])

    def test_le_reglage_de_ladministrateur_prime(self):
        os.environ["NOTIFICATION_EMAILS"] = "serveur@mairie.fr"
        try:
            c = self.connecte()
            self.assertIn("Réglage enregistré", self.poster(c, "/admin/parametres", notifications="a@mairie.fr\nb@mairie.fr, c@mairie.fr"))
            self.cree()
            self.assertEqual([m["destinataire"] for m in self.notifs()], ["a@mairie.fr", "b@mairie.fr", "c@mairie.fr"])
            self.assertIn("réglage ci-dessus", c.get("/admin/parametres").get_data(as_text=True))
            # réglage vidé : retour à la configuration du serveur
            self.assertIn("Réglage supprimé", self.poster(c, "/admin/parametres", notifications=" "))
            self.sql("DELETE FROM mails")
            self.cree(enfant_prenom="Hugo")
            self.assertEqual([m["destinataire"] for m in self.notifs()], ["serveur@mairie.fr"])
            self.assertIn("NOTIFICATION_EMAILS", c.get("/admin/parametres").get_data(as_text=True))
        finally:
            os.environ.pop("NOTIFICATION_EMAILS")

    def test_contenu_du_message(self):
        os.environ["APP_URL"] = "https://parrainage.corbas.fr/"
        self.cree(enfant_prenom="Léa", enfant_nom="Martin", parent_nom="Julie Martin", email="julie@example.org", email2="julie@example.org", essence1="2", essence2=None)
        m = self.notifs()[0]
        ref = self.un("SELECT ref FROM parrainages")["ref"]
        self.assertEqual(m["sujet"], f"Nouvelle demande de parrainage {ref} à traiter")
        self.assertIn(ref, m["corps"])
        self.assertIn("Essence souhaitée : Tilleul", m["corps"])
        self.assertIn("Demandes à traiter en ce moment : 1", m["corps"])
        self.assertIn("https://parrainage.corbas.fr/admin/demande/1", m["corps"])
        for donnee in ("Léa", "Martin", "Julie", "julie@example.org", "06 12 34 56 78", "Lilas"):
            self.assertNotIn(donnee, m["corps"] + m["sujet"])  # aucune donnée personnelle dans la notification
        self.cree(enfant_prenom="Hugo")
        self.assertIn("Demandes à traiter en ce moment : 2", self.notifs()[-1]["corps"])
        # sans APP_URL, le lien est construit à partir de l'adresse de la requête
        os.environ.pop("APP_URL")
        self.cree(enfant_prenom="Zoé")
        self.assertIn("http://localhost/admin/demande/3", self.notifs()[-1]["corps"])

    def test_liste_dattente(self):
        self.sql("UPDATE arbres SET statut='parraine'")
        self.cree(essence1="1")
        m = self.notifs()[0]
        self.assertTrue(m["sujet"].endswith("(liste d'attente)"))
        self.assertIn("rien à traiter tout de suite", m["corps"])

    def test_aucune_adresse_aucun_courrier(self):
        self.sql("DELETE FROM users")
        self.cree()
        self.assertEqual(self.notifs(), [])

    def test_page_de_reglage(self):
        self.assertEqual(self.connecte("agent@test.local").get("/admin/parametres").status_code, 403)  # réservé aux administrateurs
        c = self.connecte()
        h = c.get("/admin/parametres").get_data(as_text=True)
        self.assertIn("aucun réglage : tous les comptes", h)
        self.assertIn("admin@test.local", h)
        self.assertIn("Notifications", h)  # entrée du menu
        self.assertIn("Adresse(s) incorrecte(s) : pas-une-adresse", self.poster(c, "/admin/parametres", notifications="ok@mairie.fr\npas-une-adresse"))
        self.assertEqual(self.sql("SELECT * FROM parametres"), [])  # rien d'enregistré si une adresse est incorrecte
        trop = "\n".join(f"agent{i}@mairie.fr" for i in range(21))
        self.assertIn("20 adresses au maximum", self.poster(c, "/admin/parametres", notifications=trop))
        self.poster(c, "/admin/parametres", notifications="un@mairie.fr\nDeux@Mairie.fr")
        self.assertEqual(self.un("SELECT valeur FROM parametres WHERE cle='notification_emails'")["valeur"], "un@mairie.fr\ndeux@mairie.fr")

    def test_message_de_test(self):
        c = self.connecte()
        self.poster(c, "/admin/parametres", notifications="un@mairie.fr, deux@mairie.fr")
        h = self.poster(c, "/admin/parametres", action="test")
        self.assertIn("Message de test envoyé à 2 adresse(s)", h)
        self.assertIn("seulement enregistré", h)  # pas de SMTP dans les tests
        tests = self.sql("SELECT destinataire FROM mails WHERE sujet LIKE 'Test :%' ORDER BY id")
        self.assertEqual([m["destinataire"] for m in tests], ["un@mairie.fr", "deux@mairie.fr"])


class TestMarqueEtPhotos(Base):
    """Nom du dispositif « Mon arbre à Corbas », logo, photos d'ambiance et photos d'essences."""

    def png(self, taille=(1600, 900), couleur=(40, 120, 60)):
        from PIL import Image
        b = io.BytesIO()
        Image.new("RGB", taille, couleur).save(b, "PNG")
        return b.getvalue()

    def test_nom_du_dispositif(self):
        self.assertEqual((app.DISPOSITIF["nom"], app.DISPOSITIF["devise"]), ("Mon arbre à Corbas", "Un enfant, un arbre"))
        h = self.pub.get("/").get_data(as_text=True)
        self.assertIn("<title>Mon arbre à Corbas – Parrainez un arbre pour votre enfant</title>", h)
        self.assertIn('alt="Mon arbre à Corbas – Un enfant, un arbre"', h)
        self.assertIn("/static/mon-arbre-a-corbas.svg", h)
        self.assertIn('rel="icon"', h)
        self.assertIn("dispositif « Mon arbre à Corbas »", h)  # mention RGPD
        self.assertNotIn("🌳", h)
        login = app.app.test_client().get("/admin/login").get_data(as_text=True)
        self.assertIn("mon-arbre-a-corbas.svg", login)
        self.assertIn("Mon arbre à Corbas</b>", self.connecte().get("/admin/").get_data(as_text=True))
        # courriers : signature et premier paragraphe
        self.cree()
        accuse = self.un("SELECT corps FROM mails WHERE sujet LIKE 'Bienvenue%'")["corps"]
        self.assertIn("votre demande « Mon arbre à Corbas »", accuse)
        self.assertIn("Mon arbre à Corbas – " + app.MAIRIE["nom"], accuse)
        self.poster(self.connecte(), "/admin/parametres", action="test")
        self.assertEqual(self.un("SELECT sujet FROM mails WHERE sujet LIKE 'Test :%'")["sujet"], "Test : notifications « Mon arbre à Corbas »")

    def test_fichiers_du_logo_et_polices_livres(self):
        for chemin, type_ in (("mon-arbre-a-corbas.svg", "image/svg+xml"), ("marque.svg", "image/svg+xml"), ("marque-blanche.svg", "image/svg+xml")):
            r = self.pub.get("/static/" + chemin)
            self.assertEqual((r.status_code, r.mimetype), (200, type_), chemin)
            self.assertIn(b"<svg", r.data)
        # lettres vectorisées : le logo ne dépend d'aucune police installée
        self.assertNotIn(b"<text", self.pub.get("/static/mon-arbre-a-corbas.svg").data)
        for poids in (400, 700):
            r = self.pub.get(f"/static/fonts/roboto-{poids}.woff")
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.data[:4], b"wOFF")
        self.assertIn("fonts/roboto-400.woff", self.pub.get("/static/style.css").get_data(as_text=True))
        # les polices et images sont servies par le site lui-même : la politique de sécurité n'a rien à ouvrir
        self.assertIn("default-src 'self'", self.pub.get("/").headers["Content-Security-Policy"])

    def test_logo_du_dispositif_sur_le_certificat(self):
        self.assertTrue(os.path.isfile(app.LOGO_DISPOSITIF))
        self.assertIsNotNone(certificat.charger_logo(app.LOGO_DISPOSITIF, 56))
        self.cree()
        c = self.connecte()
        self.agir(c, 1, "valider", arbre_id=self.un("SELECT id FROM arbres WHERE code='COR-2026-001'")["id"])
        avec = c.get("/admin/demande/1/certificat.pdf").data
        fourni = app.LOGO_DISPOSITIF
        app.LOGO_DISPOSITIF = os.path.join(app.DATA, "absent.png")
        try:
            sans = c.get("/admin/demande/1/certificat.pdf").data  # sans le fichier, le certificat reste correct
        finally:
            app.LOGO_DISPOSITIF = fourni
        self.assertTrue(avec.startswith(b"%PDF") and sans.startswith(b"%PDF"))
        self.assertGreater(len(avec), len(sans))

    def test_photos_d_ambiance_facultatives(self):
        h = self.pub.get("/").get_data(as_text=True)
        if not os.path.isfile(os.path.join(app.PHOTOS, "hero-accueil.jpg")):
            self.assertIn("sans-photo", h)  # sans photo, seul le logo s'affiche : la page reste propre
            self.assertNotIn('class="hero-photo"', h)
        ancien = app.PHOTOS
        dossier = os.path.join(app.DATA, "ambiance")
        os.makedirs(dossier, exist_ok=True)
        app.PHOTOS = dossier
        try:
            h = self.pub.get("/").get_data(as_text=True)
            self.assertIn("sans-photo", h)
            self.assertNotIn('class="hero-photo"', h)
            self.assertNotIn("Crédits photos", h)
            open(os.path.join(dossier, "hero-accueil.jpg"), "wb").write(self.png())
            credit = lambda nom, auteur: {"photo": nom, "auteur": auteur, "licence": "CC BY 4.0",
                                          "licence_url": "https://creativecommons.org/licenses/by/4.0/", "source": "https://commons.wikimedia.org/wiki/File:Test.jpg"}
            with open(os.path.join(dossier, "credits.json"), "w", encoding="utf-8") as fh:
                json.dump({"hero-accueil": credit("Allée de test", "Une Personne"), "hero-merci": credit("Fleurs de test", "Une Autre")}, fh)
            h = self.pub.get("/").get_data(as_text=True)
            self.assertIn('class="hero-photo" src="/static/photos/hero-accueil.jpg"', h)
            self.assertNotIn("sans-photo", h)
            self.assertIn("Crédits photos", h)
            self.assertIn("Allée de test : Une Personne", h)
            self.assertIn('href="https://creativecommons.org/licenses/by/4.0/"', h)
            self.assertNotIn("Une Autre", h)  # seule la mention de la photo affichée figure sur la page
            # page de confirmation : bandeau court, sans photo tant que le fichier manque, donc sans mention
            merci = self.pub.get(self.envoyer().headers["Location"]).get_data(as_text=True)
            self.assertIn("hero court sans-photo", merci)
            self.assertNotIn("Crédits photos", merci)
            open(os.path.join(dossier, "hero-merci.jpg"), "wb").write(self.png())
            merci = self.pub.get(self.envoyer(enfant_prenom="Zoé").headers["Location"]).get_data(as_text=True)
            self.assertIn("Fleurs de test : Une Autre", merci)
            self.assertNotIn("Une Personne", merci)
        finally:
            app.PHOTOS = ancien
            shutil.rmtree(dossier, True)

    def test_photo_d_essence_livree_avec_lapplication(self):
        ancien = app.PHOTOS
        dossier = os.path.join(app.DATA, "ambiance")
        os.makedirs(os.path.join(dossier, "essences"), exist_ok=True)
        app.PHOTOS = dossier
        try:
            self.assertNotIn("carte-photo", self.pub.get("/").get_data(as_text=True))
            open(os.path.join(dossier, "essences", "chene.jpg"), "wb").write(self.png())
            h = self.pub.get("/").get_data(as_text=True)
            self.assertEqual(h.count('class="carte-photo"'), 1)  # seul le chêne a une photo
            self.assertIn('src="/static/photos/essences/chene.jpg"', h)
            self.assertEqual(h.count("carte avec-photo"), 1)
        finally:
            app.PHOTOS = ancien
            shutil.rmtree(dossier, True)
        self.assertEqual(app.slug("Tilleul à petites feuilles"), "tilleul-a-petites-feuilles")
        self.assertEqual(app.slug("  Chêne vert (Quercus) "), "chene-vert-quercus")

    def test_photo_d_essence_envoyee_par_ladministrateur(self):
        from PIL import Image
        c = self.connecte()
        h = self.poster(c, "/admin/essences", id="1", nom="Chêne", latin="Quercus", couleur="#2e6b34", actif="1",
                        photo=(io.BytesIO(self.png((2400, 1200))), "mon chene.png"))
        self.assertIn("Essence enregistrée", h)
        fichier = os.path.join(app.PHOTOS_ENVOYEES, "essence-1.jpg")
        with Image.open(fichier) as img:  # fermée aussitôt : Windows refuse de supprimer un fichier resté ouvert
            format_, taille = img.format, img.size
        self.assertEqual(format_, "JPEG")  # ré-encodée : le contenu envoyé n'est jamais servi tel quel
        self.assertLessEqual(max(taille), 1000)
        page = self.pub.get("/").get_data(as_text=True)
        m = re.search(r'class="carte-photo" src="(/photos/essence-1\.jpg\?v=\d+)"', page)
        self.assertTrue(m, page[:500])
        r = self.pub.get(m.group(1))
        self.assertEqual((r.status_code, r.mimetype, r.data[:2]), (200, "image/jpeg", b"\xff\xd8"))
        self.assertEqual(self.pub.get("/photos/essence-2.jpg").status_code, 404)  # pas de photo pour le tilleul
        self.assertEqual(page.count('class="carte-photo"'), 1)
        # la fiche de l'essence montre la photo et propose de la retirer
        fiche = c.get("/admin/essences?edit=1").get_data(as_text=True)
        self.assertIn("Retirer cette photo", fiche)
        # un fichier qui n'est pas une image est refusé, sans rien modifier
        refus = self.poster(c, "/admin/essences", id="1", nom="Autre nom", couleur="#2e6b34", photo=(io.BytesIO(b"<?php echo 1; ?>"), "x.jpg"))
        self.assertIn("pas reconnue", refus)
        self.assertEqual(self.un("SELECT nom FROM essences WHERE id=1")["nom"], "Chêne")
        # retrait
        self.poster(c, "/admin/essences", id="1", nom="Chêne", couleur="#2e6b34", actif="1", photo_retirer="1")
        self.assertFalse(os.path.exists(fichier))
        self.assertNotIn("carte-photo", self.pub.get("/").get_data(as_text=True))
        # nouvelle essence avec photo
        self.poster(c, "/admin/essences", nom="Hêtre", couleur="#112233", actif="1", photo=(io.BytesIO(self.png((300, 200))), "h.png"))
        eid = self.un("SELECT id FROM essences WHERE nom='Hêtre'")["id"]
        self.assertTrue(os.path.exists(os.path.join(app.PHOTOS_ENVOYEES, f"essence-{eid}.jpg")))
        # réservé aux administrateurs
        agent = self.connecte("agent@test.local")
        self.assertEqual(agent.post("/admin/essences", data=dict(csrf=self.jeton(agent), nom="X")).status_code, 403)


class TestPhotosAdmin(Base):
    """Page « Photos » : l'administrateur remplace les bandeaux du formulaire."""

    def png(self, taille=(3000, 3000), couleur=(200, 30, 30)):
        from PIL import Image
        b = io.BytesIO()
        Image.new("RGB", taille, couleur).save(b, "PNG")
        return b.getvalue()

    def test_remplacer_les_photos_d_ambiance(self):
        from PIL import Image
        agent = self.connecte("agent@test.local")
        self.assertEqual(agent.get("/admin/photos").status_code, 403)
        c = self.connecte()
        page = c.get("/admin/photos").get_data(as_text=True)
        self.assertIn("Accueil du formulaire", page)
        self.assertIn("Page de confirmation", page)
        self.assertIn('href="/admin/photos"', page)  # entrée du menu
        # au départ : la photo livrée avec l'application, avec sa mention de crédit
        accueil = self.pub.get("/").get_data(as_text=True)
        self.assertIn('class="hero-photo" src="/static/photos/hero-accueil.jpg"', accueil)
        self.assertIn("Zinneke", accueil)
        # remplacement : recadrée en 1200 x 400 ; la mention de l'ancienne photo disparaît
        self.assertIn("Photo remplacée", self.poster(c, "/admin/photos", nom="hero-accueil", photo=(io.BytesIO(self.png()), "maison.png")))
        with Image.open(os.path.join(app.PHOTOS_ENVOYEES, "hero-accueil.jpg")) as img:  # fermée aussitôt (Windows)
            format_, taille = img.format, img.size
        self.assertEqual((format_, taille), ("JPEG", (1200, 400)))
        accueil = self.pub.get("/").get_data(as_text=True)
        self.assertRegex(accueil, r'class="hero-photo" src="/photos/hero-accueil\.jpg\?v=\d+"')
        self.assertNotIn("Zinneke", accueil)
        r = self.pub.get("/photos/hero-accueil.jpg")
        self.assertEqual((r.status_code, r.mimetype), (200, "image/jpeg"))
        self.assertIn("Photo envoyée par la mairie", c.get("/admin/photos").get_data(as_text=True))
        # un fichier qui n'est pas une image, ou démesuré, est refusé ; l'autre bandeau reste intact
        self.assertIn("pas reconnue", self.poster(c, "/admin/photos", nom="hero-merci", photo=(io.BytesIO(b"pas une image"), "a.jpg")))
        enorme = io.BytesIO()
        Image.new("L", (8000, 8000)).save(enorme, "PNG")  # 64 millions de points pour quelques Ko
        self.assertLess(len(enorme.getvalue()), app.TAILLE_MAX)
        enorme.seek(0)
        self.assertIn("pas reconnue", self.poster(c, "/admin/photos", nom="hero-merci", photo=(enorme, "enorme.png")))
        self.assertIn("Choisissez la photo", self.poster(c, "/admin/photos", nom="hero-merci"))
        self.assertFalse(os.path.exists(os.path.join(app.PHOTOS_ENVOYEES, "hero-merci.jpg")))
        # seuls les deux bandeaux sont modifiables, et aucun nom ne sort du dossier
        r = c.post("/admin/photos", data=dict(csrf=self.jeton(c), nom="../secret", photo=(io.BytesIO(self.png((50, 50))), "x.png")))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.pub.get("/photos/secret.jpg").status_code, 404)
        self.assertEqual(self.pub.get("/photos/..%2Fsecret.jpg").status_code, 404)
        self.assertEqual(self.pub.get("/photos/hero-merci.jpg").status_code, 404)
        # retour à la photo d'origine
        self.assertIn("Photo d'origine rétablie", self.poster(c, "/admin/photos", nom="hero-accueil", action="origine"))
        self.assertFalse(os.path.exists(os.path.join(app.PHOTOS_ENVOYEES, "hero-accueil.jpg")))
        accueil = self.pub.get("/").get_data(as_text=True)
        self.assertIn('src="/static/photos/hero-accueil.jpg"', accueil)
        self.assertIn("Zinneke", accueil)

    def test_mentions_des_photos_d_essences_livrees(self):
        # une essence qui a une photo livrée affiche sa mention ; une photo envoyée par la mairie n'en a pas
        self.sql("UPDATE essences SET nom='Chêne vert' WHERE id=1")
        h = self.pub.get("/").get_data(as_text=True)
        self.assertIn('src="/static/photos/essences/chene-vert.jpg"', h)
        self.assertIn("Chêne vert, Montes de Vitoria : Basotxerri", h)
        c = self.connecte()
        self.poster(c, "/admin/essences", id="1", nom="Chêne vert", couleur="#2e6b34", actif="1", photo=(io.BytesIO(self.png((600, 400))), "c.png"))
        h = self.pub.get("/").get_data(as_text=True)
        self.assertNotIn("Basotxerri", h)
        self.assertIn("/photos/essence-1.jpg?v=", h)


class TestAuthentification(Base):
    def test_acces_protege(self):
        r = self.pub.get("/admin/")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/admin/login", r.headers["Location"])
        r = self.pub.post("/admin/demande/1/action", data={"action": "valider"})
        self.assertEqual(r.status_code, 302)  # session expirée : retour à la connexion, pas une page d'erreur
        self.assertIn("/admin/login", r.headers["Location"])

    def test_connexion(self):
        c = app.app.test_client()
        r = c.post("/admin/login", data=dict(email="admin@test.local", password="faux"))
        self.assertEqual(r.status_code, 200)
        self.assertIn("incorrect", r.get_data(as_text=True))
        r = c.post("/admin/login", data=dict(email="inconnu@test.local", password=MDP))
        self.assertIn("incorrect", r.get_data(as_text=True))
        r = c.post("/admin/login", data=dict(email="ADMIN@test.local", password=MDP))  # casse de l'adresse ignorée
        self.assertEqual(r.status_code, 302)
        self.assertEqual(c.get("/admin/").status_code, 200)
        c.get("/admin/logout")
        self.assertEqual(c.get("/admin/").status_code, 302)

    def test_redirection_apres_connexion(self):
        c = app.app.test_client()
        r = c.post("/admin/login?next=//evil.example", data=dict(email="admin@test.local", password=MDP))
        self.assertTrue(r.headers["Location"].endswith("/admin/"))
        c = app.app.test_client()
        r = c.post("/admin/login?next=/admin/campagnes", data=dict(email="admin@test.local", password=MDP))
        self.assertTrue(r.headers["Location"].endswith("/admin/campagnes"))
        c = app.app.test_client()
        r = c.post("/admin/login?next=https://evil.example", data=dict(email="admin@test.local", password=MDP))
        self.assertTrue(r.headers["Location"].endswith("/admin/"))

    def test_force_brute(self):
        c = app.app.test_client()
        for _ in range(10):
            c.post("/admin/login", data=dict(email="admin@test.local", password="faux"))
        r = c.post("/admin/login", data=dict(email="admin@test.local", password=MDP))
        self.assertEqual(r.status_code, 429)

    def test_csrf(self):
        self.cree()
        c = self.connecte()
        r = c.post("/admin/demande/1/action", data=dict(action="notes", notes="x"))
        self.assertEqual(r.status_code, 400)
        r = c.post("/admin/demande/1/action", data=dict(action="notes", notes="x", csrf="é"))  # jeton non ASCII : pas d'erreur 500
        self.assertEqual(r.status_code, 400)

    def test_roles(self):
        self.cree()
        ag = self.connecte("agent@test.local")
        self.assertEqual(ag.get("/admin/").status_code, 200)
        for url in ("/admin/campagnes", "/admin/essences", "/admin/arbres", "/admin/utilisateurs"):
            self.assertEqual(ag.get(url).status_code, 403, url)
        self.agir(ag, 1, "anonymiser")
        self.assertEqual(self.un("SELECT anonymise FROM parrainages")["anonymise"], 0)
        self.assertNotIn("Anonymiser", ag.get("/admin/demande/1").get_data(as_text=True))
        self.assertIn("Utilisateurs", self.connecte().get("/admin/").get_data(as_text=True))

    def test_compte_supprime_perd_sa_session(self):
        ag = self.connecte("agent@test.local")
        self.assertEqual(ag.get("/admin/").status_code, 200)
        self.sql("DELETE FROM users WHERE email='agent@test.local'")
        self.assertEqual(ag.get("/admin/").status_code, 302)

    def test_role_rafraichi(self):
        ag = self.connecte("agent@test.local")
        self.sql("UPDATE users SET role='admin' WHERE email='agent@test.local'")
        self.assertEqual(ag.get("/admin/campagnes").status_code, 200)
        self.sql("UPDATE users SET role='agent' WHERE email='agent@test.local'")
        self.assertEqual(ag.get("/admin/campagnes").status_code, 403)

    def test_gestion_des_utilisateurs(self):
        c = self.connecte()
        h = self.poster(c, "/admin/utilisateurs", action="creer", email="Nouveau@Test.local", role="agent", mdp="courtmdp")
        self.assertIn("au moins 10", h)
        h = self.poster(c, "/admin/utilisateurs", action="creer", email="nouveau@test.local", role="agent", mdp="un-bon-mdp-123")
        self.assertIn("Compte créé", h)
        h = self.poster(c, "/admin/utilisateurs", action="creer", email="NOUVEAU@test.local", role="agent", mdp="un-bon-mdp-123")
        self.assertIn("existe déjà", h)
        h = self.poster(c, "/admin/utilisateurs", action="creer", email="mauvais", role="agent", mdp="un-bon-mdp-123")
        self.assertIn("incorrecte", h)
        self.assertEqual(self.connecte("nouveau@test.local", "un-bon-mdp-123").get("/admin/").status_code, 200)
        uid = self.un("SELECT id FROM users WHERE email='nouveau@test.local'")["id"]
        self.poster(c, "/admin/utilisateurs", action="mdp", id=uid, mdp="autre-mdp-12345")
        self.connecte("nouveau@test.local", "autre-mdp-12345")
        moi = self.un("SELECT id FROM users WHERE email='admin@test.local'")["id"]
        self.assertIn("propre compte", self.poster(c, "/admin/utilisateurs", action="supprimer", id=moi))
        self.sql("DELETE FROM users WHERE email='agent@test.local'")
        self.poster(c, "/admin/utilisateurs", action="supprimer", id=uid)
        self.assertIsNone(self.un("SELECT id FROM users WHERE id=?", uid))

    def test_changement_de_mot_de_passe(self):
        c = self.connecte()
        self.assertIn("incorrect", self.poster(c, "/admin/mot-de-passe", ancien="faux", nouveau="nouveau-mdp-123", confirmation="nouveau-mdp-123"))
        self.assertIn("au moins 10", self.poster(c, "/admin/mot-de-passe", ancien=MDP, nouveau="court", confirmation="court"))
        self.assertIn("pas identiques", self.poster(c, "/admin/mot-de-passe", ancien=MDP, nouveau="nouveau-mdp-123", confirmation="autre-mdp-1234"))
        self.assertIn("modifié", self.poster(c, "/admin/mot-de-passe", ancien=MDP, nouveau="nouveau-mdp-123", confirmation="nouveau-mdp-123"))
        self.connecte("admin@test.local", "nouveau-mdp-123")
        self.assertEqual(app.app.test_client().post("/admin/login", data=dict(email="admin@test.local", password=MDP)).status_code, 200)

    def test_mot_de_passe_initial_aleatoire(self):
        dossier = tempfile.mkdtemp(dir=TMP)
        ancien = app.DB_PATH
        app.DB_PATH = os.path.join(dossier, "x.db")
        os.environ.pop("ADMIN_PASSWORD")
        try:
            app.init_db()
            self.assertEqual(sqlite3.connect(app.DB_PATH).execute("SELECT COUNT(*) FROM users").fetchone()[0], 1)
            self.assertFalse(check_password_hash(sqlite3.connect(app.DB_PATH).execute("SELECT hash FROM users").fetchone()[0], "changeme"))
        finally:
            os.environ["ADMIN_PASSWORD"] = MDP
            app.DB_PATH = ancien

    def test_cle_secrete_persistante(self):
        sauve = os.environ.pop("SECRET_KEY", None)
        chemin = os.path.join(app.DATA, "secret.key")
        if os.path.exists(chemin):
            os.remove(chemin)
        try:
            k1 = app.cle_secrete()
            self.assertEqual(len(k1), 64)
            self.assertEqual(app.cle_secrete(), k1)
            os.environ["SECRET_KEY"] = "imposee"
            self.assertEqual(app.cle_secrete(), "imposee")
        finally:
            os.environ.pop("SECRET_KEY", None)
            if sauve:
                os.environ["SECRET_KEY"] = sauve


class TestTraitement(Base):
    def setUp(self):
        super().setUp()
        self.c = self.connecte()

    def statut(self, pid=1):
        return self.un("SELECT statut FROM parrainages WHERE id=?", pid)["statut"]

    def arbre(self, code):
        return self.un("SELECT * FROM arbres WHERE code=?", code)

    def test_tableau_et_recherche(self):
        self.cree(enfant_prenom="Léa", parent_nom="Julie Martin")
        self.cree(enfant_prenom="Hugo", enfant_nom="Durand", parent_nom="Paul Durand", email="paul@example.org", email2="paul@example.org")
        h = self.c.get("/admin/").get_data(as_text=True)
        self.assertIn("Léa", h)
        self.assertIn("Hugo", h)
        h = self.c.get("/admin/?q=lea").get_data(as_text=True)  # sans accent ni casse
        self.assertIn("Léa Martin", h)
        self.assertNotIn("Hugo Durand", h)
        self.assertIn("Hugo Durand", self.c.get("/admin/?q=PAUL@example").get_data(as_text=True))
        self.assertIn("pour ces critères", self.c.get("/admin/?q=zzz").get_data(as_text=True))
        self.assertEqual(self.c.get("/admin/?statut=Remis&campagne=1").status_code, 200)

    def test_filtres_du_tableau(self):
        self.cree(enfant_prenom="Alice")
        self.cree(enfant_prenom="Bruno")
        self.cree(enfant_prenom="Carole")
        self.agir(self.c, 2, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        self.agir(self.c, 3, "complement", motif="Pièce manquante")
        h = self.c.get("/admin/?statut=a_traiter").get_data(as_text=True)  # reçues + compléments, pas les validées
        self.assertIn("Alice", h)
        self.assertIn("Carole", h)
        self.assertNotIn("Bruno", h)
        self.assertIn("2 demandes", h)
        self.assertLess(h.index("Alice"), h.index("Carole"))  # file d'attente : la plus ancienne en premier
        self.assertIn("dans l'ordre d'arrivée", h)
        tous = self.c.get("/admin/").get_data(as_text=True)
        self.assertLess(tous.index("Carole"), tous.index("Alice"))  # liste générale : la plus récente en premier
        self.assertIn("Effacer les filtres", h)
        self.assertNotIn("Effacer les filtres", self.c.get("/admin/").get_data(as_text=True))
        # copies de livret à supprimer : validée, copie encore présente
        h = self.c.get("/admin/?livret=1").get_data(as_text=True)
        self.assertIn("Bruno", h)
        self.assertNotIn("Alice", h)
        self.assertIn("copies de livret à supprimer", self.c.get("/admin/").get_data(as_text=True))
        self.agir(self.c, 2, "supprimer_livret")
        self.assertNotIn("copies de livret à supprimer", self.c.get("/admin/").get_data(as_text=True))
        self.assertIn("0 demande", self.c.get("/admin/?livret=1").get_data(as_text=True))

    def test_arbre_preselectionne_et_alerte(self):
        self.cree()
        h = self.c.get("/admin/demande/1").get_data(as_text=True)
        self.assertRegex(h, r'<option value="\d+" selected>COR-2026-001 – Chêne')  # le premier choix de la famille est proposé d'emblée
        self.assertNotIn("Valider quand même", h)
        self.sql("UPDATE parrainages SET naissance=?", (app.aujourdhui() - timedelta(days=900)).isoformat())
        self.assertIn("Valider quand même", self.c.get("/admin/demande/1").get_data(as_text=True))  # confirmation si un contrôle est en alerte
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-004")["id"])
        h = self.c.get("/admin/demande/1").get_data(as_text=True)
        self.assertRegex(h, r'<option value="\d+" selected>COR-2026-004')  # arbre déjà affecté : c'est lui qui est sélectionné

    def test_dates_lisibles(self):
        self.assertEqual(app.frdt("2026-10-29 00:08:54"), "29/10/2026 à 00h08")
        self.assertEqual((app.frdt(None), app.frdt("")), ("–", "–"))
        self.assertEqual(app.frdt("pas une date"), "pas une date")
        self.cree()
        h = self.c.get("/admin/demande/1").get_data(as_text=True)
        self.assertRegex(h, r"\d\d/\d\d/\d{4} à \d\dh\d\d")
        self.assertIn('aria-label="Avancement de la demande"', h)

    def test_validation_et_changement_darbre(self):
        self.cree()
        h = self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        self.assertIn("Demande validée", h)
        self.assertEqual(self.statut(), "Validée")
        self.assertEqual(self.arbre("COR-2026-001")["statut"], "parraine")
        mails = self.sql("SELECT * FROM mails WHERE sujet LIKE '%est prêt%'")
        self.assertEqual(len(mails), 1)
        self.assertIn("Chêne", mails[0]["corps"])
        self.assertIn("planté à l'automne 2026", mails[0]["corps"])
        # même arbre : rien ne change, pas de second courrier
        self.assertIn("déjà affecté", self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"]))
        self.assertEqual(len(self.sql("SELECT * FROM mails WHERE sujet LIKE '%est prêt%'")), 1)
        # autre arbre : l'ancien est libéré
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-004")["id"])
        self.assertEqual(self.arbre("COR-2026-001")["statut"], "libre")
        self.assertEqual(self.arbre("COR-2026-004")["statut"], "parraine")
        self.assertEqual(len(self.sql("SELECT * FROM mails WHERE sujet LIKE '%est prêt%'")), 2)

    def test_arbre_deja_pris(self):
        self.cree()
        self.cree(enfant_prenom="Hugo")
        a = self.arbre("COR-2026-001")["id"]
        self.agir(self.c, 1, "valider", arbre_id=a)
        h = self.agir(self.c, 2, "valider", arbre_id=a)
        self.assertIn("n'est plus disponible", h)
        self.assertEqual(self.statut(2), "Reçue")
        self.assertIsNone(self.un("SELECT arbre_id FROM parrainages WHERE id=2")["arbre_id"])
        self.assertEqual(self.un("SELECT COUNT(*) n FROM parrainages WHERE arbre_id=?", a)["n"], 1)
        self.assertIn("Choisissez un arbre", self.agir(self.c, 2, "valider", arbre_id=""))
        self.assertIn("Choisissez un arbre", self.agir(self.c, 2, "valider", arbre_id="9999"))

    def test_machine_a_etats(self):
        self.cree()
        self.assertIn("pas possible", self.agir(self.c, 1, "remis"))  # pas de remise sans certificat
        self.assertEqual(self.statut(), "Reçue")
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        self.assertIn("pas possible", self.agir(self.c, 1, "complement", motif="x"))
        self.assertEqual(self.statut(), "Validée")
        self.assertIn("pas possible", self.agir(self.c, 1, "remis"))
        self.c.get("/admin/demande/1/certificat.pdf")
        self.assertIn("marqué comme remis", self.agir(self.c, 1, "remis"))
        self.assertEqual(self.statut(), "Remis")
        for a in ("valider", "refuser", "attente", "complement", "remis"):
            self.assertIn("pas possible", self.agir(self.c, 1, a, arbre_id=self.arbre("COR-2026-002")["id"], motif="x"), a)
        self.assertEqual(self.statut(), "Remis")
        self.assertEqual(self.c.post("/admin/demande/1/action", data=dict(csrf=self.jeton(self.c), action="inconnue")).status_code, 400)

    def test_refus_complement_attente_annulation(self):
        self.cree()
        self.cree(enfant_prenom="Hugo")
        self.cree(enfant_prenom="Zoé")
        # complément : motif obligatoire, courrier envoyé, l'agent peut ensuite valider
        self.assertIn("Précisez", self.agir(self.c, 1, "complement", motif=""))
        self.agir(self.c, 1, "complement", motif="Page du livret illisible")
        self.assertEqual(self.statut(1), "Complément demandé")
        self.assertTrue(self.sql("SELECT 1 FROM mails WHERE corps LIKE '%illisible%'"))
        # refus : motif obligatoire, arbre libéré, copie du livret en corbeille
        self.assertIn("motif du refus est obligatoire", self.agir(self.c, 2, "refuser", motif=" "))
        self.agir(self.c, 2, "valider", arbre_id=self.arbre("COR-2026-002")["id"])
        self.agir(self.c, 2, "refuser", motif="Enfant non domicilié")
        p = self.un("SELECT * FROM parrainages WHERE id=2")
        self.assertEqual((p["statut"], p["arbre_id"], p["motif"]), ("Refusée", None, "Enfant non domicilié"))
        self.assertEqual(self.arbre("COR-2026-002")["statut"], "libre")
        self.assertIsNotNone(p["livret_supprime_at"])
        self.assertTrue(os.path.exists(os.path.join(app.LIVRETS, p["livret"])))  # récupérable 6 mois
        # liste d'attente depuis « Validée » : l'arbre revient en stock
        self.agir(self.c, 3, "valider", arbre_id=self.arbre("COR-2026-003")["id"])
        self.agir(self.c, 3, "attente")
        p = self.un("SELECT * FROM parrainages WHERE id=3")
        self.assertEqual((p["statut"], p["arbre_id"], p["valide_at"]), ("Liste d'attente", None, None))
        self.assertEqual(self.arbre("COR-2026-003")["statut"], "libre")
        # annulation : arbre libéré + livret en corbeille
        self.agir(self.c, 3, "valider", arbre_id=self.arbre("COR-2026-003")["id"])
        self.agir(self.c, 3, "annuler")
        p = self.un("SELECT * FROM parrainages WHERE id=3")
        self.assertEqual((p["statut"], p["arbre_id"]), ("Annulée", None))
        self.assertEqual(self.arbre("COR-2026-003")["statut"], "libre")
        self.assertIsNotNone(p["livret_supprime_at"])
        self.assertIn("clôturée", self.c.get("/admin/demande/3").get_data(as_text=True))

    def test_liste_attente_servie_par_une_autre_campagne(self):
        self.sql("UPDATE arbres SET statut='parraine'")
        self.cree(essence1="1")
        self.assertEqual(self.statut(), "Liste d'attente")
        self.sql("INSERT INTO campagnes(id,nom,annee,ouverte) VALUES(2,'Campagne 2027',2027,1)")
        self.sql("INSERT INTO arbres(code,campagne_id,essence_id,lieu,lat,lon,annee_plantation) VALUES('COR-2027-001',2,1,'Parc',45.67,4.9,2027)")
        h = self.c.get("/admin/demande/1").get_data(as_text=True)
        self.assertIn("COR-2027-001", h)
        self.assertIn("Campagne 2027", h)
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2027-001")["id"])
        self.assertEqual(self.statut(), "Validée")
        self.assertIn("automne 2027", self.sql("SELECT corps FROM mails WHERE sujet LIKE '%est prêt%'")[0]["corps"])

    def test_fiche_tous_statuts(self):
        self.cree()
        a = self.arbre("COR-2026-001")["id"]
        for statut in app.STATUTS:
            self.sql("UPDATE parrainages SET statut=?, arbre_id=?", statut, a if statut in ("Validée", "Certificat imprimé", "Remis") else None)
            r = self.c.get("/admin/demande/1")
            self.assertEqual(r.status_code, 200, statut)
        self.sql("UPDATE parrainages SET anonymise=1, enfant_nom='', email=''")
        self.assertIn("Fiche anonymisée", self.c.get("/admin/demande/1").get_data(as_text=True))
        self.assertEqual(self.c.get("/admin/demande/9999").status_code, 404)

    def test_controles_automatiques(self):
        self.cree()
        h = self.c.get("/admin/demande/1").get_data(as_text=True)
        self.assertIn("Dans le délai de 12 mois", h)
        self.sql("UPDATE parrainages SET naissance=?", (app.aujourdhui() - timedelta(days=900)).isoformat())
        self.assertIn("Naissance de plus de 12 mois", self.c.get("/admin/demande/1").get_data(as_text=True))
        self.sql("UPDATE parrainages SET naissance=?", (app.aujourdhui() - timedelta(days=30)).isoformat())
        self.cree(enfant_prenom="Léa", enfant_nom="Autre")  # même e-mail et même prénom : doublon possible signalé
        self.assertIn("Doublon possible : PAR-", self.c.get("/admin/demande/2").get_data(as_text=True))

    def test_notes(self):
        self.cree()
        self.agir(self.c, 1, "notes", notes="Appeler la famille")
        self.assertEqual(self.un("SELECT notes FROM parrainages")["notes"], "Appeler la famille")

    def test_certificat(self):
        self.cree()
        self.assertEqual(self.c.get("/admin/demande/1/certificat.pdf").status_code, 409)
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        r = self.c.get("/admin/demande/1/certificat.pdf")
        self.assertEqual((r.status_code, r.mimetype), (200, "application/pdf"))
        self.assertTrue(r.data.startswith(b"%PDF"))
        p = self.un("SELECT * FROM parrainages")
        self.assertEqual(p["statut"], "Certificat imprimé")
        self.assertTrue(p["imprime_at"])
        self.sql("UPDATE parrainages SET imprime_at='2026-01-02 10:00:00'")
        self.c.get("/admin/demande/1/certificat.pdf")  # une réimpression ne change ni la date ni le statut
        p = self.un("SELECT * FROM parrainages")
        self.assertEqual((p["statut"], p["imprime_at"]), ("Certificat imprimé", "2026-01-02 10:00:00"))
        self.assertIn("Marquer comme remis", self.c.get("/admin/demande/1").get_data(as_text=True))
        self.agir(self.c, 1, "remis")
        self.assertEqual(self.statut(), "Remis")
        self.assertEqual(self.c.get("/admin/demande/1/certificat.pdf").status_code, 200)  # réimpression possible après remise

    def test_certificat_cas_limites(self):
        long = " ".join(["Un texte très long sur l'essence."] * 60)
        self.sql("UPDATE essences SET symbole=?, caracteristiques=?, role=?, couleur='rouge-pas-hexa' WHERE id=1", long, long, long)
        self.cree(enfant_prenom="Łukasz Ősz-Çağla")
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        r = self.c.get("/admin/demande/1/certificat.pdf")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data.startswith(b"%PDF"))
        self.assertEqual(certificat.texte_pdf("Łukasz Ősz-Çağla œuf"), "Lukasz Osz-Çagla œuf")
        self.assertEqual(certificat.texte_pdf("a\r\nb"), "a\nb")
        self.assertEqual(certificat.texte_pdf(None), "")

    def test_logo_et_maire_sur_le_certificat(self):
        from PIL import Image, ImageDraw
        self.assertEqual(app.MAIRIE["maire"], "Benoit Eraclas")
        # le logo officiel de la ville est livré avec l'application
        self.assertTrue(os.path.isfile(app.LOGO_FOURNI))
        self.assertEqual(app.chemin_logo(), app.LOGO_FOURNI)
        r = self.pub.get("/logo")
        self.assertEqual((r.status_code, r.mimetype), (200, "image/png"))
        self.assertIn('class="logo-ville"', self.pub.get("/").get_data(as_text=True))  # visible hors iframe
        self.assertIn('class="logo-ville"', app.app.test_client().get("/admin/login").get_data(as_text=True))
        self.cree()
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        avec = self.c.get("/admin/demande/1/certificat.pdf").data
        self.assertTrue(avec.startswith(b"%PDF"))
        # sans logo, le certificat reste correct (nom de la ville en texte) et le PDF est plus léger
        fourni = app.LOGO_FOURNI
        app.LOGO_FOURNI = os.path.join(app.DATA, "absent.png")
        try:
            self.assertIsNone(app.chemin_logo())
            self.assertEqual(self.pub.get("/logo").status_code, 404)
            self.assertNotIn('class="logo-ville"', self.pub.get("/").get_data(as_text=True))
            sans = self.c.get("/admin/demande/1/certificat.pdf").data
            self.assertTrue(sans.startswith(b"%PDF"))
            self.assertGreater(len(avec), len(sans))
            # un fichier déposé dans le dossier de données remplace le logo livré
            img = Image.new("RGBA", (800, 380), (255, 255, 255, 0))
            ImageDraw.Draw(img).rectangle([10, 10, 790, 370], outline=(0, 91, 148, 255), width=20)
            chemin = os.path.join(app.DATA, "logo.png")
            img.save(chemin)
            try:
                app.LOGO_FOURNI = fourni
                self.assertEqual(app.chemin_logo(), chemin)
                open(chemin, "wb").write(b"fichier abime")  # logo illisible : le certificat reste imprimable
                self.assertTrue(self.c.get("/admin/demande/1/certificat.pdf").data.startswith(b"%PDF"))
            finally:
                os.remove(chemin)
        finally:
            app.LOGO_FOURNI = fourni
        self.assertEqual(app.chemin_logo(), fourni)

    def test_tuiles_corrompues_remplacees(self):
        dossier = tempfile.mkdtemp(dir=TMP)
        chemin = os.path.join(dossier, f"{certificat.ZOOM}_1_2.png")
        with open(chemin, "wb") as f:
            f.write(b"pas une image")
        with self.assertRaises(OSError):
            certificat._fetch_tile(1, 2, dossier)  # hors ligne : la tuile tronquée est supprimée puis le téléchargement échoue
        self.assertFalse(os.path.exists(chemin))

    def test_livret_corbeille_et_restauration(self):
        self.cree()
        ag = self.connecte("agent@test.local")
        self.assertEqual(ag.get("/admin/demande/1/livret").status_code, 200)
        self.assertEqual(ag.get("/admin/demande/1/livret").headers["Cache-Control"], "no-store")
        self.assertIn("Copie du livret supprimée", self.agir(ag, 1, "supprimer_livret"))
        self.assertEqual(ag.get("/admin/demande/1/livret").status_code, 403)  # l'agent n'y a plus accès
        self.assertEqual(self.c.get("/admin/demande/1/livret").status_code, 200)  # l'administrateur, si
        t1 = self.un("SELECT livret_supprime_at FROM parrainages")["livret_supprime_at"]
        self.agir(ag, 1, "supprimer_livret")  # un second clic ne repousse pas l'échéance
        self.assertEqual(self.un("SELECT livret_supprime_at FROM parrainages")["livret_supprime_at"], t1)
        self.agir(ag, 1, "restaurer_livret")  # réservé à l'administrateur
        self.assertIsNotNone(self.un("SELECT livret_supprime_at FROM parrainages")["livret_supprime_at"])
        self.assertIn("restaurée", self.agir(self.c, 1, "restaurer_livret"))
        self.assertIsNone(self.un("SELECT livret_supprime_at FROM parrainages")["livret_supprime_at"])
        self.assertEqual(ag.get("/admin/demande/1/livret").status_code, 200)

    def test_purge_apres_six_mois(self):
        self.cree()
        self.cree(enfant_prenom="Hugo")
        fichier = lambda i: os.path.join(app.LIVRETS, self.un("SELECT livret FROM parrainages WHERE id=?", i)["livret"])
        f1, f2 = fichier(1), fichier(2)
        vieux = (app.maintenant() - timedelta(days=app.CONSERVATION_JOURS + 1)).strftime("%Y-%m-%d %H:%M:%S")
        recent = (app.maintenant() - timedelta(days=app.CONSERVATION_JOURS - 2)).strftime("%Y-%m-%d %H:%M:%S")
        self.sql("UPDATE parrainages SET livret_supprime_at=? WHERE id=1", vieux)
        self.sql("UPDATE parrainages SET livret_supprime_at=? WHERE id=2", recent)
        with app.app.app_context():
            app.purger_livrets()
        self.assertFalse(os.path.exists(f1))
        self.assertTrue(os.path.exists(f2))
        self.assertIsNone(self.un("SELECT livret FROM parrainages WHERE id=1")["livret"])
        self.assertIsNotNone(self.un("SELECT livret_supprime_at FROM parrainages WHERE id=1")["livret_supprime_at"])
        self.assertIn("effacée définitivement", self.c.get("/admin/demande/1").get_data(as_text=True))
        self.assertIn("Il n'y a plus de copie", self.agir(self.c, 1, "restaurer_livret"))
        self.assertEqual(self.c.get("/admin/demande/1/livret").status_code, 404)

    def test_purge_tolerante_aux_fichiers_verrouilles(self):
        self.cree()
        self.sql("UPDATE parrainages SET livret_supprime_at=?", (app.maintenant() - timedelta(days=400)).strftime("%Y-%m-%d %H:%M:%S"))
        vrai = os.remove
        os.remove = lambda *a: (_ for _ in ()).throw(PermissionError("verrouillé"))
        try:
            with app.app.app_context():
                app.purger_livrets()  # ne plante pas
        finally:
            os.remove = vrai
        self.assertIsNotNone(self.un("SELECT livret FROM parrainages")["livret"])  # référence conservée : nouvel essai à la prochaine purge
        with app.app.app_context():
            app.purger_livrets()
        self.assertIsNone(self.un("SELECT livret FROM parrainages")["livret"])

    def test_anonymisation(self):
        self.cree()
        self.cree(enfant_prenom="Hugo", enfant_nom="Durand", parent_nom="Paul Durand", email="paul@example.org", email2="paul@example.org")
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        f1 = os.path.join(app.LIVRETS, self.un("SELECT livret FROM parrainages WHERE id=1")["livret"])
        self.assertIn("Données personnelles effacées", self.agir(self.c, 1, "anonymiser"))
        p = self.un("SELECT * FROM parrainages WHERE id=1")
        self.assertEqual((p["enfant_nom"], p["email"], p["tel"], p["adresse"], p["parent_nom"], p["anonymise"]), ("", "", "", "", "(anonymisé)", 1))
        self.assertIsNone(p["livret"])
        self.assertFalse(os.path.exists(f1))
        self.assertEqual(p["arbre_id"], self.arbre("COR-2026-001")["id"])  # la fiche de l'arbre est conservée
        tous = " ".join(f"{m['destinataire']} {m['sujet']} {m['corps']}" for m in self.sql("SELECT * FROM mails"))
        self.assertNotIn("julie@example.org", tous)
        self.assertNotIn("Julie Martin", tous)
        self.assertNotIn("Léa", tous)
        self.assertTrue(self.sql("SELECT 1 FROM mails WHERE destinataire='paul@example.org'"))  # l'autre famille n'est pas touchée
        self.assertEqual(self.c.post("/admin/demande/1/action", data=dict(csrf=self.jeton(self.c), action="notes", notes="x")).status_code, 400)

    def test_motif_efface_par_anonymisation(self):
        self.cree()
        self.agir(self.c, 1, "complement", motif="Votre adresse 3 rue des Lilas est incomplète")
        self.agir(self.c, 1, "anonymiser")
        self.assertIsNone(self.un("SELECT motif FROM parrainages")["motif"])
        self.assertFalse(self.sql("SELECT 1 FROM mails WHERE corps LIKE '%Lilas%'"))

    def test_export_excel(self):
        self.cree(enfant_prenom="=1+1", enfant_nom="@SUM(1)", parent_nom="+cmd|calc", adresse='=HYPERLINK("http://evil","clic")')
        self.agir(self.c, 1, "valider", arbre_id=self.arbre("COR-2026-001")["id"])
        r = self.c.get("/admin/export.xlsx")
        self.assertEqual(r.status_code, 200)
        z = zipfile.ZipFile(io.BytesIO(r.data))
        for nom in z.namelist():
            if nom.startswith("xl/worksheets/"):
                self.assertNotIn("<f>", z.read(nom).decode("utf-8"), nom)  # aucune formule : tout est du texte
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ["Parrainages", "Arbres", "Essences", "Campagnes"])
        ws = wb["Parrainages"]
        self.assertEqual(ws["E2"].value, "=1+1")
        self.assertEqual(ws["E2"].data_type, "s")
        self.assertEqual(ws["T2"].value, "COR-2026-001")
        self.assertEqual(wb["Arbres"].max_row, 7)
        self.assertEqual(wb["Arbres"]["L2"].value, ws["A2"].value)
        self.assertEqual(wb["Essences"]["A2"].value, "Chêne")

    def test_courriers_sans_smtp(self):
        self.cree()
        h = self.c.get("/admin/mails").get_data(as_text=True)
        self.assertIn("enregistré seulement", h)
        self.assertNotIn("Renvoyer", h)
        self.assertIn("Aucun serveur", self.poster(self.c, "/admin/mails/1/renvoyer"))

    def test_envoi_smtp_en_arriere_plan(self):
        envoyes, vrai = [], app._smtp

        def faux(to, sujet, corps):
            if "echec" in to:
                raise ConnectionError("serveur injoignable")
            envoyes.append(to)
        app._smtp = faux
        os.environ["SMTP_HOST"] = "smtp.test"
        try:
            self.cree()
            self.cree(enfant_prenom="Hugo", email="echec@example.org", email2="echec@example.org")
            for _ in range(60):
                if len(self.sql("SELECT 1 FROM mails WHERE envoye=1")) >= 3 and self.sql("SELECT 1 FROM mails WHERE erreur IS NOT NULL"):
                    break
                time.sleep(0.1)
            self.assertIn("julie@example.org", envoyes)
            ko = self.sql("SELECT * FROM mails WHERE destinataire='echec@example.org'")[0]
            self.assertEqual((ko["envoye"], ko["erreur"]), (0, "serveur injoignable"))
            h = self.c.get("/admin/mails").get_data(as_text=True)
            self.assertIn("échec : serveur injoignable", h)
            self.assertIn("Renvoyer", h)
            # le serveur revient : le courrier est renvoyé
            app._smtp = lambda to, s, c: envoyes.append("renvoi:" + to)
            self.assertIn("Nouvel envoi", self.poster(self.c, f"/admin/mails/{ko['id']}/renvoyer"))
            for _ in range(60):
                if self.un("SELECT envoye FROM mails WHERE id=?", ko["id"])["envoye"]:
                    break
                time.sleep(0.1)
            self.assertEqual(self.un("SELECT envoye, erreur FROM mails WHERE id=?", ko["id"])["envoye"], 1)
            self.assertIn("renvoi:echec@example.org", envoyes)
        finally:
            app._smtp = vrai
            os.environ.pop("SMTP_HOST", None)


class TestParametrage(Base):
    def setUp(self):
        super().setUp()
        self.c = self.connecte()

    def test_campagnes(self):
        for kw, msg in ((dict(nom=" ", annee="2027"), "nom"), (dict(nom="X", annee="1900"), "année"),
                        (dict(nom="X", annee="2027", ouverture="2027-05-01", fermeture="2027-04-01"), "précéder"),
                        (dict(nom="X", annee="2027", ouverture="pas-une-date"), "dates")):
            self.assertIn(msg, self.poster(self.c, "/admin/campagnes", **kw), kw)
        self.assertEqual(self.un("SELECT COUNT(*) n FROM campagnes")["n"], 1)
        self.assertIn("Campagne enregistrée", self.poster(self.c, "/admin/campagnes", nom="Printemps 2027", annee="2027", ouvert="1", ouverte="1",
                                                          ouverture="2027-03-01", fermeture="2027-05-31"))
        cid = self.un("SELECT id FROM campagnes WHERE nom='Printemps 2027'")["id"]
        self.poster(self.c, "/admin/campagnes", id=cid, nom="Printemps 2027 bis", annee="2027")
        self.assertEqual(self.un("SELECT nom, ouverte FROM campagnes WHERE id=?", cid)["ouverte"], 0)
        self.assertEqual(self.c.get(f"/admin/campagnes?edit={cid}").status_code, 200)

    def test_essences(self):
        self.assertIn("nom à l'essence", self.poster(self.c, "/admin/essences", nom=""))
        self.assertIn("couleur doit s'écrire", self.poster(self.c, "/admin/essences", nom="Hêtre", couleur="vert"))
        self.assertIn("Essence enregistrée", self.poster(self.c, "/admin/essences", nom="Hêtre", latin="Fagus", couleur="#112233", actif="1"))
        self.assertEqual(self.un("SELECT couleur, actif FROM essences WHERE nom='Hêtre'")["couleur"], "#112233")
        self.assertEqual(self.c.get("/admin/essences?edit=1").status_code, 200)

    def test_arbres_ajout_et_codes(self):
        h = self.poster(self.c, "/admin/arbres", campagne_id="1", essence_id="1", nombre="3", lieu="Parc Bourlione", adresse="", lat="45,6706", lon="4.8960", annee_plantation="2027")
        self.assertIn("3 arbre(s) ajouté(s)", h)
        codes = [r["code"] for r in self.sql("SELECT code FROM arbres ORDER BY code")]
        self.assertEqual(codes[-3:], ["COR-2026-007", "COR-2026-008", "COR-2026-009"])
        self.assertEqual(self.un("SELECT lat FROM arbres WHERE code='COR-2026-007'")["lat"], 45.6706)
        # après la suppression d'un arbre du milieu, le code suivant ne doit pas entrer en collision
        self.sql("DELETE FROM arbres WHERE code='COR-2026-003'")
        self.poster(self.c, "/admin/arbres", campagne_id="1", essence_id="2", nombre="1", lieu="Square", lat="45.67", lon="4.9")
        self.assertIsNotNone(self.un("SELECT 1 FROM arbres WHERE code='COR-2026-010'"))
        self.assertEqual(self.un("SELECT COUNT(*) n FROM arbres")["n"], 9)

    def test_arbres_saisies_invalides(self):
        base = dict(campagne_id="1", essence_id="1", nombre="1", lieu="Parc", lat="45.67", lon="4.9")
        avant = self.un("SELECT COUNT(*) n FROM arbres")["n"]
        for kw, msg in ((dict(lat="nan"), "Vérifiez"), (dict(lat="inf"), "Vérifiez"), (dict(lat="abc"), "Vérifiez"), (dict(lat="48.96"), "pas à Corbas"),
                        (dict(lon="49.0"), "pas à Corbas"), (dict(nombre="0"), "Vérifiez"), (dict(nombre="201"), "Vérifiez"),
                        (dict(nombre="x"), "Vérifiez"), (dict(lieu="  "), "Vérifiez"), (dict(campagne_id="99"), "Vérifiez"),
                        (dict(essence_id="99"), "Vérifiez")):
            self.assertIn(msg, self.poster(self.c, "/admin/arbres", **{**base, **kw}), kw)
        self.assertEqual(self.un("SELECT COUNT(*) n FROM arbres")["n"], avant)

    def test_arbres_statut_et_suppression(self):
        self.cree()
        a1 = self.un("SELECT id FROM arbres WHERE code='COR-2026-001'")["id"]
        a2 = self.un("SELECT id FROM arbres WHERE code='COR-2026-002'")["id"]
        self.agir(self.c, 1, "valider", arbre_id=a1)
        # un arbre affecté ne peut pas repasser « libre » (il serait proposé à une seconde famille)
        self.assertIn("annulez d'abord", self.poster(self.c, "/admin/arbres", action="statut", id=a1, statut="libre"))
        self.assertEqual(self.un("SELECT statut FROM arbres WHERE id=?", a1)["statut"], "parraine")
        self.assertIn("validation d'une demande", self.poster(self.c, "/admin/arbres", action="statut", id=a2, statut="parraine"))
        self.assertIn("inconnu", self.poster(self.c, "/admin/arbres", action="statut", id=a2, statut="nimportequoi"))
        self.assertIn("incorrecte", self.poster(self.c, "/admin/arbres", action="statut", id=a2, statut="libre", date_plantation="demain"))
        self.assertIn("mis à jour", self.poster(self.c, "/admin/arbres", action="statut", id=a1, statut="remplace", date_plantation="2026-11-20"))
        self.assertEqual(self.un("SELECT date_plantation FROM arbres WHERE id=?", a1)["date_plantation"], "2026-11-20")
        self.assertIn("mis à jour", self.poster(self.c, "/admin/arbres", action="statut", id=a2, statut="libre", date_plantation="2026-11-21"))
        self.assertIn("Seul un arbre libre", self.poster(self.c, "/admin/arbres", action="supprimer", id=a1))
        self.assertIn("supprimé du stock", self.poster(self.c, "/admin/arbres", action="supprimer", id=a2))
        self.assertEqual(self.c.get("/admin/arbres?campagne=1").status_code, 200)
        self.assertEqual(self.c.get("/admin/arbres").status_code, 200)


if __name__ == "__main__":
    unittest.main(verbosity=2)
